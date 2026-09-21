#!/usr/bin/env python3
"""Build the deterministic, redistributable CoFID evaluation projection."""

from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "FoodSourceEvaluation"))
from analyse_cofid import CATALOGUE_MAPPING, EXPECTED_SHA256, UNMAPPED_CATALOGUE_KEYS  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


SOURCE_RELEASE_ID = f"cofid:2021:sha256:{EXPECTED_SHA256}"


def canonical_json(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n").encode()


def canonical_decimal(value: Any) -> str | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        decimal = Decimal(str(value).strip())
    except InvalidOperation:
        return None
    rendered = format(decimal.normalize(), "f")
    return "0" if rendered in {"-0", ""} else rendered


def nutrient_value(value: Any, source_unit: str, canonical_unit: str) -> dict[str, Any]:
    units = {"source_unit": source_unit, "canonical_unit": canonical_unit}
    if value is None or str(value).strip() == "":
        return {"state": "blank", **units}
    token = str(value).strip()
    if token.casefold() == "tr":
        return {"state": "trace", **units}
    if token.casefold() == "n":
        return {"state": "present_unquantified", **units}
    numeric = canonical_decimal(value)
    if numeric is not None:
        return {"state": "numeric", **units, "value": numeric}
    return {"state": "other", **units, "source_token": token}


def contains_any(name: str, values: tuple[str, ...]) -> bool:
    return any(value in name for value in values)


def inferred_identity(name: str, published_group: str) -> dict[str, str]:
    lowered = name.casefold()
    if "raw" in lowered:
        preparation = "raw"
    elif contains_any(lowered, ("cooked", "boiled", "fried", "grilled", "roasted", "baked", "stewed", "steamed")):
        preparation = "cooked"
    else:
        preparation = "unknown"
    if contains_any(lowered, ("with bone", "weighed with bone", "bone in", "bone-in")):
        bone = "with_bone"
    elif "boneless" in lowered:
        bone = "boneless"
    else:
        bone = "unknown"
    if contains_any(lowered, ("skin on", "skin-on", "with skin")):
        skin = "skin_on"
    elif contains_any(lowered, ("skinless", "without skin")):
        skin = "skinless"
    else:
        skin = "unknown"
    drained = "undrained" if "undrained" in lowered else "drained" if "drained" in lowered else "unknown"
    if "olive oil" in lowered:
        packing_medium = "olive_oil"
    elif "brine" in lowered:
        packing_medium = "brine"
    elif "in oil" in lowered:
        packing_medium = "oil"
    elif "in water" in lowered:
        packing_medium = "water"
    elif "sauce" in lowered:
        packing_medium = "sauce"
    else:
        packing_medium = "unknown"
    fortification = "unfortified" if "unfortified" in lowered else "fortified" if "fortified" in lowered else "unknown"
    salt_state = "unsalted" if "unsalted" in lowered else "salted" if "salted" in lowered else "unknown"
    return {
        "preparation": preparation,
        "bone": bone,
        "skin": skin,
        "drained": drained,
        "packing_medium": packing_medium,
        "fortification": fortification,
        "salt_state": salt_state,
        "serving_basis": "per_100_ml" if published_group.startswith("Q") else "per_100_g",
        "edible_quantity": "unknown",
        "formulation": "unknown",
    }


def build_release(workbook_path: Path) -> dict[str, Any]:
    workbook_bytes = workbook_path.read_bytes()
    digest = hashlib.sha256(workbook_bytes).hexdigest()
    if digest != EXPECTED_SHA256:
        raise ValueError(f"unexpected CoFID SHA-256: {digest}; expected {EXPECTED_SHA256}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheets = {
        name: list(workbook[name].iter_rows(values_only=True))
        for name in sorted({mapping[0] for mapping in CATALOGUE_MAPPING.values()})
    }
    primary = sheets["1.3 Proximates"]
    indexes: dict[tuple[str, str], int] = {}
    for sheet_name, rows in sheets.items():
        indexes.update(
            {
                (sheet_name, str(code).strip()): index
                for index, code in enumerate(rows[1])
                if code is not None
            }
        )

    records = []
    for worksheet_row, row in enumerate(primary[3:], start=4):
        if row[0] is None:
            continue
        code = str(row[0]).strip()
        name = str(row[1]).strip()
        description = str(row[2]).strip() if row[2] is not None else ""
        published_group = str(row[3]).strip() if row[3] is not None else ""
        row_identity = hashlib.sha256(canonical_json({"row": worksheet_row, "code": code, "name": name})).hexdigest()
        nutrients = {}
        for key, (sheet_name, source_code, canonical_unit) in CATALOGUE_MAPPING.items():
            sheet_row = sheets[sheet_name][worksheet_row - 1]
            if str(sheet_row[0]).strip() != code or str(sheet_row[1]).strip() != name:
                raise ValueError(
                    f"CoFID row alignment drift at worksheet row {worksheet_row} on {sheet_name}"
                )
            source_unit = "g" if key == "water" else canonical_unit
            nutrients[key] = nutrient_value(
                sheet_row[indexes[(sheet_name, source_code)]], source_unit, canonical_unit
            )
        for key in UNMAPPED_CATALOGUE_KEYS:
            canonical_unit = {"chromium": "mcg", "molybdenum": "mcg", "caffeine": "mg"}[key]
            nutrients[key] = {
                "state": "unmapped",
                "source_unit": None,
                "canonical_unit": canonical_unit,
            }
        records.append(
            {
                "record_id": f"{SOURCE_RELEASE_ID}:row:{worksheet_row:04d}:{row_identity[:16]}",
                "source_release_id": SOURCE_RELEASE_ID,
                "worksheet_row": worksheet_row,
                "published_code": code,
                "published_group": published_group,
                "name": name,
                "description": description,
                "identity": inferred_identity(name, published_group),
                "nutrients": nutrients,
            }
        )
    return {
        "schema_version": 1,
        "source": {
            "source_id": "cofid",
            "release_id": SOURCE_RELEASE_ID,
            "title": "McCance and Widdowson's Composition of Foods Integrated Dataset 2021",
            "artifact_sha256": digest,
            "licence": "Open Government Licence v3.0",
            "source_url": "https://www.gov.uk/government/publications/composition-of-foods-integrated-dataset-cofid",
            "attribution": "Contains public sector information licensed under the Open Government Licence v3.0.",
            "record_identity": "worksheet row plus canonical row hash; published food codes are not unique",
        },
        "records": records,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    payload = canonical_json(build_release(args.workbook))
    args.output.write_bytes(gzip.compress(payload, compresslevel=9, mtime=0))
    print(json.dumps({"records": len(json.loads(payload)["records"]), "canonical_sha256": hashlib.sha256(payload).hexdigest(), "gzip_sha256": hashlib.sha256(args.output.read_bytes()).hexdigest()}))


if __name__ == "__main__":
    main()
