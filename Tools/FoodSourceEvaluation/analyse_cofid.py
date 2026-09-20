#!/usr/bin/env python3
"""Measure the CoFID 2021 workbook against the repository nutrient contract."""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


CATALOGUE_MAPPING = {
    "energy_consumed": ("1.3 Proximates", "KCALS", "kcal"),
    "carbohydrates": ("1.3 Proximates", "CHO", "g"),
    "protein": ("1.3 Proximates", "PROT", "g"),
    "fat_total": ("1.3 Proximates", "FAT", "g"),
    "fat_saturated": ("1.3 Proximates", "SATFOD", "g"),
    "fat_monounsaturated": ("1.3 Proximates", "MONOFOD", "g"),
    "fat_polyunsaturated": ("1.3 Proximates", "POLYFOD", "g"),
    "fiber": ("1.3 Proximates", "AOACFIB", "g"),
    "sugar": ("1.3 Proximates", "TOTSUG", "g"),
    "cholesterol": ("1.3 Proximates", "CHOL", "mg"),
    "vitamin_a": ("1.5 Vitamins", "RETEQU", "mcg"),
    "thiamin_b1": ("1.5 Vitamins", "THIA", "mg"),
    "riboflavin_b2": ("1.5 Vitamins", "RIBO", "mg"),
    "niacin_b3": ("1.5 Vitamins", "NIAC", "mg"),
    "pantothenic_acid_b5": ("1.5 Vitamins", "PANTO", "mg"),
    "vitamin_b6": ("1.5 Vitamins", "VITB6", "mg"),
    "biotin_b7": ("1.5 Vitamins", "BIOT", "mcg"),
    "folate_b9": ("1.5 Vitamins", "FOLT", "mcg"),
    "vitamin_b12": ("1.5 Vitamins", "VITB12", "mcg"),
    "vitamin_c": ("1.5 Vitamins", "VITC", "mg"),
    "vitamin_d": ("1.5 Vitamins", "VITD", "mcg"),
    "vitamin_e": ("1.5 Vitamins", "VITE", "mg"),
    "vitamin_k": ("1.5 Vitamins", "VITK1", "mcg"),
    "calcium": ("1.4 Inorganics", "CA", "mg"),
    "chloride": ("1.4 Inorganics", "CL", "mg"),
    "copper": ("1.4 Inorganics", "CU", "mg"),
    "iodine": ("1.4 Inorganics", "I", "mcg"),
    "iron": ("1.4 Inorganics", "FE", "mg"),
    "magnesium": ("1.4 Inorganics", "MG", "mg"),
    "manganese": ("1.4 Inorganics", "MN", "mg"),
    "phosphorus": ("1.4 Inorganics", "P", "mg"),
    "potassium": ("1.4 Inorganics", "K", "mg"),
    "selenium": ("1.4 Inorganics", "SE", "mcg"),
    "sodium": ("1.4 Inorganics", "NA", "mg"),
    "zinc": ("1.4 Inorganics", "ZN", "mg"),
    "water": ("1.3 Proximates", "WATER", "mL"),
}

UNMAPPED_CATALOGUE_KEYS = ["chromium", "molybdenum", "caffeine"]
EXPECTED_SHA256 = "436e9445ef2adb2a75f3d7edd51302de3adad25385f9795fc94ba58bd030e97d"
IDENTITY_NAME_MARKERS = [
    "fortified",
    "drained",
    "undrained",
    "weighed with bone",
    "boneless",
    "supplement",
    "canned",
]


def value_state(value: Any) -> str:
    if value is None or str(value).strip() == "":
        return "blank"
    text = str(value).strip().lower()
    if text == "tr":
        return "trace"
    if text == "n":
        return "present_unquantified"
    try:
        float(value)
        return "numeric"
    except (TypeError, ValueError):
        return "other"


def analyse(path: Path) -> dict[str, Any]:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_rows: dict[str, list[tuple[Any, ...]]] = {}
    for sheet_name in {entry[0] for entry in CATALOGUE_MAPPING.values()}:
        sheet_rows[sheet_name] = list(workbook[sheet_name].iter_rows(values_only=True))

    primary_rows = sheet_rows["1.3 Proximates"]
    food_codes = [str(row[0]).strip() for row in primary_rows[3:] if row[0] is not None]
    duplicate_codes = sorted(code for code, count in Counter(food_codes).items() if count > 1)
    duplicate_records = [
        {"worksheet_row": index, "food_code": str(row[0]).strip(), "food_name": row[1]}
        for index, row in enumerate(primary_rows[3:], start=4)
        if row[0] is not None and str(row[0]).strip() in duplicate_codes
    ]
    food_names = [str(row[1]).lower() for row in primary_rows[3:] if row[0] is not None]

    nutrient_results = []
    for catalogue_key, (sheet_name, source_code, canonical_unit) in CATALOGUE_MAPPING.items():
        rows = sheet_rows[sheet_name]
        source_codes = {str(value).strip(): index for index, value in enumerate(rows[1]) if value is not None}
        column = source_codes[source_code]
        states = Counter(value_state(row[column]) for row in rows[3:] if row[0] is not None)
        nutrient_results.append(
            {
                "catalogue_key": catalogue_key,
                "canonical_unit": canonical_unit,
                "source_sheet": sheet_name,
                "source_code": source_code,
                "source_heading": rows[0][column],
                "counts": {key: states.get(key, 0) for key in ("numeric", "trace", "present_unquantified", "blank", "other")},
            }
        )

    return {
        "artifact": {
            "filename": path.name,
            "bytes": path.stat().st_size,
            "sha256": digest,
            "expected_sha256_match": digest == EXPECTED_SHA256,
        },
        "release": "CoFID 2021",
        "food_rows": len(food_codes),
        "unique_food_codes": len(set(food_codes)),
        "duplicate_food_codes": duplicate_codes,
        "duplicate_records": duplicate_records,
        "food_name_marker_counts": {
            marker: sum(marker in name for name in food_names) for marker in IDENTITY_NAME_MARKERS
        },
        "contract_nutrients": 39,
        "mapped_nutrients": len(CATALOGUE_MAPPING),
        "unmapped_catalogue_keys": UNMAPPED_CATALOGUE_KEYS,
        "nutrients": nutrient_results,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = analyse(args.workbook)
    rendered = json.dumps(result, indent=2, ensure_ascii=False) + "\n"
    if args.output:
        args.output.write_text(rendered, encoding="utf-8")
    else:
        print(rendered, end="")


if __name__ == "__main__":
    main()
